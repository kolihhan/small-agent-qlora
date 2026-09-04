from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from .base import Tool, ToolResult
from .content_type import detect_content_type


class InspectTool(Tool):
    name = "inspect"
    description = "Inspect a workspace file's structure/metadata without reading all content."
    schema = {
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
        "additionalProperties": False,
    }

    def run(self, arguments: dict[str, Any], workspace: Path) -> ToolResult:
        raw = arguments.get("path")
        if not isinstance(raw, str) or not raw.strip():
            return ToolResult(False, "'path' must be a non-empty string", "BAD_ARGUMENTS")
        path = (workspace / raw).resolve()
        try:
            path.relative_to(workspace.resolve())
        except ValueError:
            return ToolResult(False, "File must be inside the workspace", "PATH_OUTSIDE_WORKSPACE")
        if not path.exists():
            return ToolResult(False, f"Not found: {raw}", "NOT_FOUND")
        stat = path.stat()
        detected_type = detect_content_type(path) if path.is_file() else "directory"
        info: dict[str, Any] = {"name": path.name, "suffix": path.suffix.lower(), "detected_type": detected_type, "bytes": stat.st_size}
        try:
            if detected_type == "xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(path.open("rb"), read_only=True, data_only=False)
                info["sheets"] = [{"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column} for ws in wb.worksheets]
            elif detected_type == "pdf":
                from pypdf import PdfReader
                reader = PdfReader(path)
                info["pages"] = len(reader.pages)
                info["metadata"] = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
            elif detected_type in {"csv", "tsv"}:
                import csv
                delim = "\t" if detected_type == "tsv" else ","
                with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
                    reader = csv.reader(f, delimiter=delim)
                    first = next(reader, [])
                    rows = 1 + sum(1 for _ in reader) if first else 0
                info.update({"rows": rows, "columns": len(first), "header": first})
        except ImportError as exc:
            info["optional_inspection"] = f"Unavailable: {exc}"
        except Exception as exc:
            info["inspection_warning"] = f"{type(exc).__name__}: {exc}"
        return ToolResult(True, json.dumps(info, ensure_ascii=False, indent=2))
